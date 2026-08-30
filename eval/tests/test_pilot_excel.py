"""Tests for the spreadsheet form of the rubric pilot.

The thing that matters here is the round trip. A workbook that cannot be read back is worse than
no workbook at all: the annotators' half-day is spent and the result is unusable. So most of
these build a real file, fill it the way a person would, and read it back.

Network-free. Everything runs against temporary files.
"""

from __future__ import annotations

import itertools
import os

import openpyxl
import pytest

from eval.gold_set import load_gold_set
from eval.judge.checklist import load_checklist
from eval.judge.pilot import arm_results, fixed_pilot_items, pilot_status
from eval.judge.pilot_excel import (
    SHEET_ANCHORED,
    SHEET_CHECKLIST,
    SHEET_ROLLUP,
    SHEET_TIME,
    ExcelPilotError,
    build_workbook,
    checklist_rows,
    estimate_workload,
    pilot_dimensions,
    read_workbook,
)
from eval.judge.rubric import load_rubric

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
GOLD_PATH = os.path.join(REPO_ROOT, "gold_set.json")

pytestmark = pytest.mark.skipif(
    not os.path.exists(GOLD_PATH), reason="gold_set.json not present in the repo root"
)

_INPUT_YELLOW = "FFF2CC"


@pytest.fixture(scope="module")
def items():
    return fixed_pilot_items(load_gold_set(GOLD_PATH))


def is_input(cell) -> bool:
    colour = cell.fill.start_color.rgb
    return bool(colour) and str(colour).endswith(_INPUT_YELLOW)


def fill(path, anchored=(4, 3, 2, 1), ticks=("Yes", "Yes", "No", "Yes"), minutes=(45, 60)):
    """Fill a workbook the way an annotator would, and report how many cells were filled."""
    workbook = openpyxl.load_workbook(path)
    counts = {}
    for sheet_name, column, values in (
        (SHEET_ANCHORED, 3, anchored), (SHEET_CHECKLIST, 3, ticks)
    ):
        sheet = workbook[sheet_name]
        cycle = itertools.cycle(values)
        filled = 0
        for row in range(1, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column)
            if is_input(cell):
                cell.value = next(cycle)
                filled += 1
        counts[sheet_name] = filled
    if minutes:
        log = workbook[SHEET_TIME]
        log.cell(row=2, column=2).value = minutes[0]
        log.cell(row=3, column=2).value = minutes[1]
    workbook.save(path)
    return counts


# ---------------------------------------------------------------------------
# What gets scored
# ---------------------------------------------------------------------------
def test_both_arms_score_the_same_dimensions(items):
    """Otherwise the arms are not comparable, which is the one thing the pilot measures."""
    checklist = load_checklist()
    rubric = load_rubric()
    for item in items:
        anchored = set(pilot_dimensions(item, rubric))
        from_checklist = {row[0] for row in checklist_rows(item, checklist, rubric)}
        assert anchored == from_checklist, item.id


def test_refusal_items_are_scored_only_on_refusal_appropriateness(items):
    for item in items:
        if item.should_refuse:
            assert pilot_dimensions(item) == ["refusal_appropriateness"]


def test_citation_validity_is_not_scored_because_the_answers_do_not_cite(items):
    """A real gap in this pilot, pinned so it is a known limitation rather than a surprise."""
    every = {d for item in items for d in pilot_dimensions(item)}
    assert "citation_validity" not in every
    assert "temporal_integrity" not in every


def test_the_workload_estimate_matches_the_cells_actually_written(items, tmp_path):
    path = str(tmp_path / "w.xlsx")
    estimate = build_workbook(items, path, annotator_id="ann_a")
    counts = fill(path)
    assert counts[SHEET_ANCHORED] == estimate.anchored_cells
    assert counts[SHEET_CHECKLIST] == estimate.tick_boxes


def test_arm_three_costs_no_annotator_time(items):
    """It is arm 2's ticks rolled up, so the estimate must not charge for it three times."""
    estimate = estimate_workload(items)
    assert estimate.minutes < 180
    assert estimate.anchored_cells == 28
    assert estimate.tick_boxes == 112


# ---------------------------------------------------------------------------
# The round trip
# ---------------------------------------------------------------------------
def test_a_filled_workbook_reads_back_as_all_three_arms(items, tmp_path):
    path = str(tmp_path / "ann_a.xlsx")
    estimate = build_workbook(items, path, annotator_id="ann_a")
    fill(path)

    labels = read_workbook(path)
    by_arm = {}
    for label in labels:
        by_arm.setdefault(label.arm, []).append(label)

    assert set(by_arm) == {"anchored", "checklist", "rollup"}
    assert len(by_arm["anchored"]) == estimate.anchored_cells
    # One checklist unit per (item, dimension), not per tick box.
    assert len(by_arm["checklist"]) == len(by_arm["rollup"]) == 28
    assert all(label.annotator_id == "ann_a" for label in labels)


def test_units_are_item_and_dimension_so_twelve_items_give_more_than_twelve_units(items, tmp_path):
    """At n=12 a per-item unit would leave kappa on twelve points. Per dimension gives 28."""
    path = str(tmp_path / "a.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    fill(path)
    anchored = [l for l in read_workbook(path) if l.arm == "anchored"]
    assert len(anchored) == 28
    assert all("::" in label.item_id for label in anchored)


def test_the_annotator_id_comes_from_the_workbook(items, tmp_path):
    path = str(tmp_path / "b.xlsx")
    build_workbook(items, path, annotator_id="ann_b")
    fill(path)
    assert {l.annotator_id for l in read_workbook(path)} == {"ann_b"}


def test_an_unnamed_workbook_is_refused_rather_than_scored_anonymously(items, tmp_path):
    path = str(tmp_path / "anon.xlsx")
    build_workbook(items, path, annotator_id="")
    fill(path)
    with pytest.raises(ExcelPilotError, match="does not name an annotator"):
        read_workbook(path)
    # An explicit id rescues it.
    assert read_workbook(path, annotator_id="ann_c")


def test_time_is_spread_across_units_not_charged_whole_to_each(items, tmp_path):
    path = str(tmp_path / "t.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    fill(path, minutes=(45, 60))
    labels = read_workbook(path)

    anchored = next(l for l in labels if l.arm == "anchored")
    assert anchored.seconds == pytest.approx(45 * 60 / 28)
    # Arm 3 inherits arm 2's time: it costs no extra effort and must not look more expensive.
    checklist = next(l for l in labels if l.arm == "checklist")
    rollup = next(l for l in labels if l.arm == "rollup")
    assert checklist.seconds == rollup.seconds


def test_a_missing_time_log_is_not_fatal(items, tmp_path):
    path = str(tmp_path / "notime.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    fill(path, minutes=None)
    assert all(label.seconds is None for label in read_workbook(path))


# ---------------------------------------------------------------------------
# Blanks and bad input
# ---------------------------------------------------------------------------
def test_blank_and_na_become_none_and_are_never_imputed(items, tmp_path):
    path = str(tmp_path / "gaps.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    fill(path, anchored=(4, "NA", 3, ""))

    anchored = [l for l in read_workbook(path) if l.arm == "anchored"]
    blanks = [l for l in anchored if l.score is None]
    assert blanks, "NA and empty must survive as None, not as a guessed middle value"
    assert all(l.score in (None, 3.0, 4.0) for l in anchored)


def test_an_out_of_range_score_is_refused(items, tmp_path):
    path = str(tmp_path / "bad.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    fill(path, anchored=(7,))
    with pytest.raises(ExcelPilotError, match="outside 1-4"):
        read_workbook(path)


def test_a_non_numeric_score_is_refused(items, tmp_path):
    path = str(tmp_path / "words.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    fill(path, anchored=("good",))
    with pytest.raises(ExcelPilotError, match="expected 1-4"):
        read_workbook(path)


def test_an_untouched_workbook_says_so(items, tmp_path):
    path = str(tmp_path / "empty.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    with pytest.raises(ExcelPilotError, match="no scores filled in"):
        read_workbook(path)


def test_a_missing_file_is_refused(tmp_path):
    with pytest.raises(ExcelPilotError, match="no workbook at"):
        read_workbook(str(tmp_path / "nope.xlsx"))


def test_building_with_no_items_is_refused(tmp_path):
    with pytest.raises(ExcelPilotError):
        build_workbook([], str(tmp_path / "x.xlsx"))


# ---------------------------------------------------------------------------
# Arm 3's arithmetic
# ---------------------------------------------------------------------------
def test_all_yes_rolls_up_to_four_and_all_no_to_one(items, tmp_path):
    top = str(tmp_path / "top.xlsx")
    build_workbook(items, top, annotator_id="ann_a")
    fill(top, ticks=("Yes",))
    assert {l.score for l in read_workbook(top) if l.arm == "rollup"} == {4.0}

    bottom = str(tmp_path / "bottom.xlsx")
    build_workbook(items, bottom, annotator_id="ann_a")
    fill(bottom, ticks=("No",))
    assert {l.score for l in read_workbook(bottom) if l.arm == "rollup"} == {1.0}


def test_na_ticks_leave_the_denominator_alone(items, tmp_path):
    """An inapplicable box should not count against the score."""
    path = str(tmp_path / "na.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    fill(path, ticks=("Yes", "NA"))
    rollup = [l for l in read_workbook(path) if l.arm == "rollup"]
    # Every scored box that was answered was a Yes, so nothing should fall below the top band.
    assert set(r.score for r in rollup) == {4.0}


def test_the_sheet_carries_a_live_formula_for_arm_three(items, tmp_path):
    """The annotator must be able to see arm 3 derive itself, not take it on trust."""
    path = str(tmp_path / "f.xlsx")
    build_workbook(items, path, annotator_id="ann_a")
    sheet = openpyxl.load_workbook(path)[SHEET_ROLLUP]
    formulas = [
        sheet.cell(row=row, column=5).value
        for row in range(3, 3 + len(items))
    ]
    assert all(isinstance(f, str) and f.startswith("=IF(") for f in formulas)
    earned = sheet.cell(row=3, column=2).value
    assert "SUMIF" in earned and SHEET_CHECKLIST in earned


# ---------------------------------------------------------------------------
# Straight into the existing scorer
# ---------------------------------------------------------------------------
def test_two_workbooks_score_through_the_unchanged_pilot_scorer(items, tmp_path):
    paths = []
    for annotator, anchored in (("ann_a", (4, 3, 2, 1)), ("ann_b", (4, 3, 2, 1))):
        path = str(tmp_path / f"{annotator}.xlsx")
        build_workbook(items, path, annotator_id=annotator)
        fill(path, anchored=anchored)
        paths.append(path)

    labels = [label for path in paths for label in read_workbook(path)]
    results = {r.arm: r for r in arm_results(labels)}

    assert set(results) == {"anchored", "checklist", "rollup"}
    # Identical answers from both annotators is perfect agreement.
    assert results["anchored"].agreement == pytest.approx(1.0)

    status = pilot_status(labels)
    assert status.passed is True
    assert status.winner in {"anchored", "checklist", "rollup"}


def test_disagreeing_annotators_lower_the_score(items, tmp_path):
    paths = []
    for annotator, anchored in (("ann_a", (4, 1, 4, 1)), ("ann_b", (1, 4, 1, 4))):
        path = str(tmp_path / f"{annotator}.xlsx")
        build_workbook(items, path, annotator_id=annotator)
        fill(path, anchored=anchored)
        paths.append(path)

    labels = [label for path in paths for label in read_workbook(path)]
    anchored_result = next(r for r in arm_results(labels) if r.arm == "anchored")
    assert anchored_result.agreement is not None
    assert anchored_result.agreement < 0.6
