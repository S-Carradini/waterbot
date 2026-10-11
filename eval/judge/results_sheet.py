"""Build the human-review workbook for a FACTUAL panel run.

Not the qualitative sheet. `qualitative_sheet.py` writes into the RealWater evaluator TEMPLATE,
blinds the judges, hides the scores, and exists so a human can rate the JUDGE's reasoning. This
one is the opposite job: a flat, sortable, filterable view of what the panel actually scored on
the 30-item gold set, so a person can skim 450 verdicts without reading raw JSON.

THE JSONL IS THE SOURCE OF TRUTH. This workbook is generated from it and is disposable - never
hand-edit it, regenerate it. Aggregation, agreement statistics and the report gate all read the
jsonl directly and depend on it being append-only (invariant 2). A score changed in a
spreadsheet cell would be invisible to every one of them.

ONE ROW PER (item x judge x dimension). 30 items x 3 judges x 5 rubric dimensions is 450 rows
at most - fewer, because a dimension only appears on the items its `applies_to` covers. The k=5
repeats are not their own rows: the five scores sit in one column as a list, with their mean
beside it, and the `remark` is the rationale from the sample closest to that judge's own median
(the same representative-sample choice `qualitative_sheet` makes), so the text shown is the
reasoning behind the judge's typical verdict rather than an outlier.
"""
from __future__ import annotations

import os
import statistics
from collections import defaultdict
from typing import Any, Dict, List, Optional, Sequence

from eval.judge.qualitative_sheet import load_verdicts, representative_sample

SHEET = "Panel Results"

HEADERS = [
    "item_id", "category", "question", "chatbot_answer",
    "judge", "dimension",
    "k scores", "mean", "abstentions",
    "remark (representative sample)", "quoted span", "evidence ids", "repaired any",
]


def _responses_by_item(path: str) -> Dict[str, Dict[str, Any]]:
    """item_id -> {question, category, answer} from a frozen responses/*.jsonl."""
    import json

    out: Dict[str, Dict[str, Any]] = {}
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            d = json.loads(line)
            out[d["item_id"]] = {
                "question": d.get("question", ""),
                "category": d.get("category", ""),
                "answer": d.get("answer", ""),
            }
    return out


def _mean(xs: Sequence[Optional[float]]) -> Optional[float]:
    real = [x for x in xs if x is not None]
    return round(sum(real) / len(real), 2) if real else None


def build(verdict_paths: Sequence[str], responses_path: str, out: str) -> Dict[str, Any]:
    """Write the review workbook. Returns a small summary dict."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    verdicts = load_verdicts(verdict_paths)
    responses = _responses_by_item(responses_path)

    by_item_judge: Dict[str, Dict[str, List[Any]]] = defaultdict(lambda: defaultdict(list))
    for v in verdicts:
        by_item_judge[v.item_id][v.judge_model_id].append(v)

    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = SHEET
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    rows_written = 0
    for item_id in sorted(by_item_judge):
        meta = responses.get(item_id, {"question": "", "category": "", "answer": ""})
        for judge in sorted(by_item_judge[item_id]):
            samples = by_item_judge[item_id][judge]
            representative = representative_sample(samples)
            # Every dimension that appears on any sample of this (item, judge).
            dims = sorted({d for v in samples for d in v.dimensions})
            for dim in dims:
                scores = [v.dimensions[dim].score for v in samples if dim in v.dimensions]
                abstentions = sum(1 for s in scores if s is None)
                rep_entry = representative.dimensions.get(dim)
                if rep_entry is None:
                    # The representative sample did not score this dimension; fall back to the
                    # first sample that did, so the remark column is never silently blank.
                    rep_entry = next(
                        (v.dimensions[dim] for v in samples if dim in v.dimensions), None)
                remark = rep_entry.rationale.strip() if rep_entry else ""
                span = (rep_entry.quoted_span or "").strip() if rep_entry else ""
                ev_ids = ", ".join(str(i) for i in rep_entry.evidence_ids) if rep_entry else ""
                ws.append([
                    item_id,
                    meta["category"],
                    meta["question"],
                    meta["answer"],
                    judge,
                    dim,
                    ", ".join("abstain" if s is None else str(s) for s in scores),
                    _mean(scores),
                    abstentions,
                    remark,
                    span,
                    ev_ids,
                    "yes" if any(v.repaired for v in samples) else "no",
                ])
                rows_written += 1

    widths = {"A": 10, "B": 22, "C": 46, "D": 60, "E": 20, "F": 22,
              "G": 16, "H": 8, "I": 12, "J": 90, "K": 50, "L": 12, "M": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    workbook.save(out)
    return {
        "rows": rows_written,
        "verdicts": len(verdicts),
        "items": len(by_item_judge),
        "judges": sorted({v.judge_model_id for v in verdicts}),
        "out": out,
    }
